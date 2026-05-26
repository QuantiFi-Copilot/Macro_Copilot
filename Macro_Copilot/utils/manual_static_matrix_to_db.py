"""Manual static matrix ingester — Bloomberg BDP fields → instrument_master.attributes.

Reads a static matrix XLSX (one sheet, columns: vendor_ticker |
instrument_type | NAME | SECURITY_DES | BB_FIRST_TRADE_DT |
CRNCY_QUOTATION_CONV | MARKET_SECTOR_DES | TICKER_AND_EXCH_CODE |
PRICE_TYPE) and MERGES the non-N/A BDP fields into the
`macro_data.instrument_master.attributes` JSONB column via the
PostgreSQL `||` operator. Never replaces the JSONB — only adds /
overrides specific keys.

This is distinct from the time-series flow (`manual_bbg_xlsx_to_parquet.py`
+ `ingestion/ingest_parquet.py`). Static matrix data has 1 scalar value
per ticker (not a date-indexed series), so it lives in `instrument_master`
attributes rather than `market_data_daily`.

Idempotent: re-running the same XLSX produces zero net effect on rows
where the attrs already match.

Dry-run by default; pass --apply to commit.

Usage:
    # Dry-run (default)
    python utils/manual_static_matrix_to_db.py \\
        --input ~/Downloads/fx_bidask_warehouse_20260525/static_matrix_extraction_20260525.xlsx

    # Commit
    python utils/manual_static_matrix_to_db.py \\
        --input ~/Downloads/fx_bidask_warehouse_20260525/static_matrix_extraction_20260525.xlsx \\
        --apply

Note: must be run with localhost:5433 reachable (the Mac-side Docker
tsdb container must be UP). For Docker-only environments, the same
SQL pattern can be invoked from inside the rates-agent-dev container.
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path
from typing import Dict, List, Tuple

import openpyxl
from sqlalchemy import create_engine, text


# Map column header → JSONB attribute key (snake_case for consistency
# with other attributes already in instrument_master).
#
# Skip columns we don't want to persist or that BBG returned #N/A for
# all FX OTC instruments (verified Phase 1 discovery 2026-05-26):
#   - instrument_type — already in instrument_master.instrument_type
#   - BB_FIRST_TRADE_DT, CRNCY_QUOTATION_CONV, TICKER_AND_EXCH_CODE,
#     PRICE_TYPE — all #N/A for FX OTC (equity-centric fields)
COL_HEADER_TO_ATTR_KEY: Dict[str, str] = {
    "NAME": "name",
    "SECURITY_DES": "security_des",
    "MARKET_SECTOR_DES": "market_sector_des",
}

# Values to treat as missing. Covers Bloomberg's various #N/A variants
# plus the obvious None/empty-string.
_NA_LITERAL_VALUES: set = {
    "",
    "#N/A N/A",
    "#N/A Invalid Field",
    "#N/A Field Not Applicable",
    "#N/A Field",
    "#N/A",
    "N/A",
}


def is_na(v) -> bool:
    """Return True if the BDP value should be treated as missing."""
    if v is None:
        return True
    if isinstance(v, str):
        # Bloomberg returns various #N/A flavors; treat all as missing
        stripped = v.strip()
        if stripped in _NA_LITERAL_VALUES:
            return True
        if stripped.startswith("#N/A"):
            return True
    return False


def read_static_matrix(xlsx_path: Path) -> List[Tuple[str, Dict[str, str]]]:
    """Read XLSX, return list of (vendor_ticker, attrs_dict) where attrs
    only contains BDP fields that resolved to a non-N/A value.

    Tickers with zero usable attrs are excluded from the result.
    """
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    sheet_name = wb.sheetnames[0]
    ws = wb[sheet_name]

    # Read header from row 1 to map column index → attr key
    header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    col_idx_to_attr_key: Dict[int, str] = {}
    for i, col_name in enumerate(header_row):
        if col_name in COL_HEADER_TO_ATTR_KEY:
            col_idx_to_attr_key[i] = COL_HEADER_TO_ATTR_KEY[col_name]

    if not col_idx_to_attr_key:
        raise ValueError(
            f"No usable columns found in static matrix XLSX. "
            f"Expected one of {list(COL_HEADER_TO_ATTR_KEY)}. "
            f"Header found: {list(header_row)}"
        )

    rows: List[Tuple[str, Dict[str, str]]] = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r or len(r) == 0:
            continue
        ticker = r[0]
        if not ticker or not isinstance(ticker, str):
            continue
        ticker = ticker.strip()
        if not ticker:
            continue
        attrs: Dict[str, str] = {}
        for col_idx, attr_key in col_idx_to_attr_key.items():
            v = r[col_idx] if col_idx < len(r) else None
            if is_na(v):
                continue
            attrs[attr_key] = str(v).strip()
        if attrs:
            rows.append((ticker, attrs))

    wb.close()
    return rows


def apply_to_db(
    rows: List[Tuple[str, Dict[str, str]]],
    *,
    dry_run: bool = True,
) -> int:
    """Apply (ticker, attrs) updates to instrument_master.attributes via
    JSONB || merge (never replace). Returns the number of rows that
    would be / were applied.
    """
    engine = create_engine(
        "postgresql://quantuser:myStrongPass@localhost:5433/macrodata"
    )

    # First pass: which tickers exist in DB?
    with engine.connect() as conn:
        existing = conn.execute(
            text(
                "SELECT vendor_ticker FROM macro_data.instrument_master "
                "WHERE vendor_ticker = ANY(:t)"
            ),
            {"t": [t for t, _ in rows]},
        ).fetchall()
        existing_tickers = {r[0] for r in existing}

    n_to_update = sum(1 for t, _ in rows if t in existing_tickers)
    n_missing = sum(1 for t, _ in rows if t not in existing_tickers)
    n_total = len(rows)

    print(f"Static matrix: {n_total} tickers in XLSX with at least 1 non-NA attr")
    print(f"  - existing in DB : {n_to_update}")
    print(f"  - missing from DB: {n_missing}")

    if n_missing > 0:
        missing_list = [t for t, _ in rows if t not in existing_tickers]
        print(f"  - missing examples (up to 10): {missing_list[:10]}")

    if n_to_update == 0:
        print("\nNothing to apply (no ticker overlap with instrument_master).")
        return 0

    if dry_run:
        # Show sample of what would be applied
        sample = [r for r in rows if r[0] in existing_tickers][:5]
        print(f"\nDry-run sample (first 5 of {n_to_update} updates):")
        for ticker, attrs in sample:
            print(f"  {ticker}:")
            for k, v in attrs.items():
                v_short = v if len(v) <= 80 else v[:77] + "..."
                print(f"    {k}: {v_short!r}")
        print(
            f"\nDry-run only. Re-run with --apply to commit "
            f"{n_to_update} updates."
        )
        return n_to_update

    # APPLY mode: execute UPDATEs in a single transaction.
    # Uses JSONB || (merge) operator — preserves all existing keys not
    # in `attrs`, overrides keys that ARE in `attrs`. Idempotent.
    sql = text(
        """
        UPDATE macro_data.instrument_master
        SET attributes = COALESCE(attributes, '{}'::jsonb) || CAST(:attrs AS jsonb)
        WHERE vendor_ticker = :ticker
        """
    )

    with engine.begin() as conn:
        n_applied = 0
        for ticker, attrs in rows:
            if ticker not in existing_tickers:
                continue
            conn.execute(sql, {"ticker": ticker, "attrs": json.dumps(attrs)})
            n_applied += 1
        print(f"\n[DB] Applied {n_applied} JSONB merge updates in 1 transaction.")
    print("[DB] Done.")
    return n_applied


def main() -> int:
    parser = argparse.ArgumentParser(
        description=__doc__,
        formatter_class=argparse.RawDescriptionHelpFormatter,
    )
    parser.add_argument(
        "--input", required=True, type=Path,
        help="Path to the static matrix XLSX (typically "
             "static_matrix_extraction_YYYYMMDD.xlsx).",
    )
    parser.add_argument(
        "--apply", action="store_true",
        help="Commit the JSONB merge updates. Default is dry-run.",
    )
    args = parser.parse_args()

    if not args.input.exists():
        print(f"ERROR: input file not found: {args.input}", file=sys.stderr)
        return 1

    print(f"Reading static matrix from {args.input} ...")
    try:
        rows = read_static_matrix(args.input)
    except Exception as exc:
        print(f"ERROR: failed to read XLSX: {exc}", file=sys.stderr)
        return 1

    if not rows:
        print("ERROR: no tickers with usable non-NA attrs.", file=sys.stderr)
        return 1

    apply_to_db(rows, dry_run=not args.apply)
    return 0


if __name__ == "__main__":
    sys.exit(main())
